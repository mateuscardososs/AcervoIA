from collections.abc import Iterator
from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document as WordDocument
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from acervo_ia import config
from acervo_ia.api.routes import documents
from acervo_ia.db.connection import get_db
from acervo_ia.db.models import Base, Collection, Document, DocumentChunk, User
from acervo_ia.main import app
from acervo_ia.security import create_access_token
from acervo_ia.services.document_processing import TextPage, chunk_pages


@pytest.fixture
def processing_client(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> Iterator[TestClient]:
    monkeypatch.setenv(
        "AUTH_SECRET_KEY",
        "test-auth-secret-key-with-sufficient-length-0123456789",
    )
    monkeypatch.setattr(
        documents,
        "DOCUMENT_STORAGE_DIRECTORY",
        tmp_path / "stored",
    )
    engine = create_engine(
        "sqlite+pysqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    test_session_factory = sessionmaker(
        bind=engine,
        autoflush=False,
        expire_on_commit=False,
    )
    previous_session_factory = getattr(app.state, "test_session_factory", None)
    app.state.test_session_factory = test_session_factory

    def override_get_db() -> Iterator[Session]:
        with test_session_factory() as session:
            yield session

    previous_override = app.dependency_overrides.get(get_db)
    app.dependency_overrides[get_db] = override_get_db

    try:
        with TestClient(app) as test_client:
            yield test_client
    finally:
        if previous_override is None:
            app.dependency_overrides.pop(get_db, None)
        else:
            app.dependency_overrides[get_db] = previous_override
        if previous_session_factory is None:
            del app.state.test_session_factory
        else:
            app.state.test_session_factory = previous_session_factory
        engine.dispose()


def pdf_bytes(text: str) -> bytes:
    content_stream = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"
        ),
        b"<< /Length "
        + str(len(content_stream)).encode()
        + b" >>\nstream\n"
        + content_stream
        + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref_offset = len(pdf)
    pdf.extend(f"xref\n0 {len(offsets)}\n".encode())
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode())
    pdf.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n".encode()
    )
    return bytes(pdf)


def docx_bytes() -> bytes:
    document = WordDocument()
    document.add_paragraph("Texto principal do DOCX")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Texto da tabela"
    content = BytesIO()
    document.save(content)
    return content.getvalue()


def create_document(
    email: str,
    filename: str,
    content_type: str,
    content: bytes,
) -> tuple[UUID, UUID, UUID, str, Path]:
    with app.state.test_session_factory() as session:
        user = User(email=email, password_hash="unused")
        session.add(user)
        session.flush()
        collection = Collection(owner_id=user.id, name=f"Coleção de {email}")
        session.add(collection)
        session.flush()
        storage_key = uuid4().hex
        stored_path = documents.DOCUMENT_STORAGE_DIRECTORY / storage_key
        stored_path.parent.mkdir(parents=True, exist_ok=True)
        stored_path.write_bytes(content)
        document = Document(
            collection_id=collection.id,
            original_filename=filename,
            storage_key=storage_key,
            content_type=content_type,
            size_bytes=len(content),
        )
        session.add(document)
        session.commit()
        user_id = user.id
        collection_id = collection.id
        document_id = document.id
    return (
        user_id,
        collection_id,
        document_id,
        create_access_token(str(user_id)),
        stored_path,
    )


@pytest.mark.parametrize(
    "filename,content_type,content,expected_text,expected_page",
    [
        (
            "manual.txt",
            "text/plain",
            "Texto simples em UTF-8.".encode(),
            "Texto simples em UTF-8.",
            None,
        ),
        (
            "manual.pdf",
            "application/pdf",
            pdf_bytes("Texto extraido do PDF."),
            "Texto extraido do PDF.",
            1,
        ),
        (
            "manual.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            docx_bytes(),
            "Texto principal do DOCX",
            None,
        ),
    ],
)
def test_processes_txt_pdf_and_docx(
    processing_client: TestClient,
    filename: str,
    content_type: str,
    content: bytes,
    expected_text: str,
    expected_page: int | None,
) -> None:
    _, collection_id, document_id, token, _ = create_document(
        "owner@example.test",
        filename,
        content_type,
        content,
    )

    response = processing_client.post(
        f"/collections/{collection_id}/documents/{document_id}/process",
        headers={"Authorization": f"Bearer {token}"},
    )

    assert response.status_code == 200
    assert response.json()["processing_status"] == "completed"
    assert response.json()["chunk_count"] == 1
    with app.state.test_session_factory() as session:
        document = session.get(Document, document_id)
        chunks = session.scalars(
            select(DocumentChunk).where(DocumentChunk.document_id == document_id)
        ).all()
    assert document is not None
    assert document.processing_error is None
    assert len(chunks) == 1
    assert expected_text in chunks[0].content
    assert chunks[0].position == 0
    assert chunks[0].page_number == expected_page
    if filename.endswith(".docx"):
        assert "Texto da tabela" in chunks[0].content


def test_chunking_applies_configured_size_and_overlap(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(config, "CHUNK_SIZE_CHARS", 6)
    monkeypatch.setattr(config, "CHUNK_OVERLAP_CHARS", 2)

    chunks = chunk_pages([TextPage(text="abcdefghij", page_number=None)])

    assert [chunk.content for chunk in chunks] == ["abcdef", "efghij"]
    assert [chunk.page_number for chunk in chunks] == [None, None]


def test_pdf_chunks_keep_page_numbers(
    processing_client: TestClient,
) -> None:
    _, collection_id, document_id, token, _ = create_document(
        "pdf@example.test",
        "manual.pdf",
        "application/pdf",
        pdf_bytes("Pagina um."),
    )

    response = processing_client.post(
        f"/collections/{collection_id}/documents/{document_id}/process",
        headers={"Authorization": f"Bearer {token}"},
    )

    assert response.status_code == 200
    with app.state.test_session_factory() as session:
        chunk = session.scalar(
            select(DocumentChunk).where(DocumentChunk.document_id == document_id)
        )
    assert chunk is not None
    assert chunk.page_number == 1


def test_reprocessing_replaces_previous_chunks(
    processing_client: TestClient,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(config, "CHUNK_SIZE_CHARS", 12)
    monkeypatch.setattr(config, "CHUNK_OVERLAP_CHARS", 3)
    _, collection_id, document_id, token, stored_path = create_document(
        "owner@example.test",
        "manual.txt",
        "text/plain",
        b"A" * 25,
    )
    endpoint = f"/collections/{collection_id}/documents/{document_id}/process"
    headers = {"Authorization": f"Bearer {token}"}

    first = processing_client.post(endpoint, headers=headers)
    stored_path.write_bytes(b"B" * 25)
    second = processing_client.post(endpoint, headers=headers)

    assert first.status_code == second.status_code == 200
    assert first.json()["chunk_count"] == second.json()["chunk_count"] == 3
    with app.state.test_session_factory() as session:
        chunks = session.scalars(
            select(DocumentChunk)
            .where(DocumentChunk.document_id == document_id)
            .order_by(DocumentChunk.position)
        ).all()
    assert len(chunks) == 3
    assert [chunk.position for chunk in chunks] == [0, 1, 2]
    assert all(set(chunk.content) == {"B"} for chunk in chunks)


def test_document_without_extractable_text_fails_safely_and_keeps_file(
    processing_client: TestClient,
) -> None:
    _, collection_id, document_id, token, stored_path = create_document(
        "empty@example.test",
        "empty.txt",
        "text/plain",
        b" \n\t ",
    )

    response = processing_client.post(
        f"/collections/{collection_id}/documents/{document_id}/process",
        headers={"Authorization": f"Bearer {token}"},
    )

    assert response.status_code == 422
    assert response.json() == {
        "detail": "Não foi possível extrair texto do documento."
    }
    assert stored_path.is_file()
    with app.state.test_session_factory() as session:
        document = session.get(Document, document_id)
        assert document is not None
        assert document.processing_status == "failed"
        assert document.processing_error == response.json()["detail"]
        assert session.scalar(
            select(DocumentChunk).where(DocumentChunk.document_id == document_id)
        ) is None


def test_processing_checks_collection_owner_before_opening_file(
    processing_client: TestClient,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _, collection_id, document_id, _, _ = create_document(
        "owner@example.test",
        "manual.txt",
        "text/plain",
        b"Private document text",
    )
    other_user = User(email="other@example.test", password_hash="unused")
    with app.state.test_session_factory() as session:
        session.add(other_user)
        session.commit()
        other_token = create_access_token(str(other_user.id))

    def file_path_must_not_be_resolved(_storage_key: str) -> Path:
        pytest.fail("O arquivo foi acessado antes de validar o dono da coleção")

    monkeypatch.setattr(documents, "_safe_storage_path", file_path_must_not_be_resolved)
    response = processing_client.post(
        f"/collections/{collection_id}/documents/{document_id}/process",
        headers={"Authorization": f"Bearer {other_token}"},
    )

    assert response.status_code == 404
